import os
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from app.schemas.cv_schema import CvDataSchema

def set_cell_margins(cell, top=50, bottom=50, left=50, right=50):
    """Establece márgenes internos en celdas de tabla"""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def add_horizontal_rule(paragraph, color_hex="000000", size_eighth_pt="12"):
    """Agrega una línea divisoria horizontal inferior nativa al párrafo (100% compatible ATS)"""
    pPr = paragraph._element.get_or_add_pPr()
    pBdr = parse_xml(
        f'<w:pBdr {nsdecls("w")}>\n'
        f'  <w:bottom w:val="single" w:sz="{size_eighth_pt}" w:space="2" w:color="{color_hex}"/>\n'
        f'</w:pBdr>'
    )
    pPr.append(pBdr)

def build_ats_docx(cv: CvDataSchema, output_path: str) -> str:
    """
    Genera un documento Word (.docx) cumpliendo estrictamente los estándares ATS:
    - 1 Sola Columna
    - Alineación a la izquierda
    - Fuente estándar Calibri de alta legibilidad
    - Viñetas sólidas nativas ('List Bullet')
    - Cero cuadros de texto flotantes, imágenes o tablas anidadas
    """
    doc = Document()

    # Configuración de márgenes estándar (0.75 pulgadas)
    for section in doc.sections:
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(0.75)
        section.right_margin = Inches(0.75)

    # Configuración de estilo Normal
    style_normal = doc.styles['Normal']
    style_normal.font.name = 'Calibri'
    style_normal.font.size = Pt(10)
    style_normal.font.color.rgb = RGBColor(30, 30, 30)

    h = cv.header

    # -----------------------------------------------------------------
    # 1. ENCABEZADO Y CONTACTO
    # -----------------------------------------------------------------
    p_name = doc.add_paragraph()
    p_name.paragraph_format.space_before = Pt(0)
    p_name.paragraph_format.space_after = Pt(2)
    run_name = p_name.add_run(h.full_name.upper())
    run_name.font.size = Pt(19)
    run_name.font.bold = True
    run_name.font.color.rgb = RGBColor(0, 0, 0)

    p_title = doc.add_paragraph()
    p_title.paragraph_format.space_before = Pt(0)
    p_title.paragraph_format.space_after = Pt(4)
    run_title = p_title.add_run(h.professional_title)
    run_title.font.size = Pt(11)
    run_title.font.bold = True
    run_title.font.color.rgb = RGBColor(37, 99, 235)  # Azul profesional discreto

    # Línea de contacto: Ubicación • Teléfono • Email
    contact_parts = []
    location = ", ".join(filter(None, [h.city, h.country]))
    if location:
        contact_parts.append(location)
    if h.phone:
        contact_parts.append(h.phone)
    if h.email:
        contact_parts.append(h.email)

    p_contact = doc.add_paragraph()
    p_contact.paragraph_format.space_before = Pt(0)
    p_contact.paragraph_format.space_after = Pt(2)
    p_contact.add_run("  •  ".join(contact_parts))

    # Línea de enlaces: LinkedIn • GitHub (opcional) • Portafolio (opcional)
    link_parts = []
    if h.linkedin_url:
        link_parts.append(f"LinkedIn: {h.linkedin_url}")
    if h.github_url:
        link_parts.append(f"GitHub: {h.github_url}")
    if h.portfolio_url:
        link_parts.append(f"Portafolio: {h.portfolio_url}")

    if link_parts:
        p_links = doc.add_paragraph()
        p_links.paragraph_format.space_before = Pt(0)
        p_links.paragraph_format.space_after = Pt(12)
        run_links = p_links.add_run("  •  ".join(link_parts))
        run_links.font.size = Pt(9.5)
        run_links.font.color.rgb = RGBColor(37, 99, 235)

    def add_section_header(title: str):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(title.upper())
        run.font.size = Pt(11)
        run.font.bold = True
        run.font.color.rgb = RGBColor(0, 0, 0)
        add_horizontal_rule(p, color_hex="000000", size_eighth_pt="8")
        return p

    # -----------------------------------------------------------------
    # 2. RESUMEN PROFESIONAL (ELEVATOR PITCH - 3 A 4 LÍNEAS)
    # -----------------------------------------------------------------
    add_section_header("Resumen Profesional")
    p_summary = doc.add_paragraph()
    p_summary.paragraph_format.space_before = Pt(3)
    p_summary.paragraph_format.space_after = Pt(8)
    p_summary.paragraph_format.line_spacing = 1.15
    run_sum = p_summary.add_run(h.summary)
    run_sum.font.size = Pt(10)

    # -----------------------------------------------------------------
    # 3. COMPETENCIAS TÉCNICAS (HARD SKILLS & HERRAMIENTAS)
    # -----------------------------------------------------------------
    if cv.skill_categories:
        add_section_header("Competencias Técnicas")
        for cat in cv.skill_categories:
            p_skill = doc.add_paragraph()
            p_skill.paragraph_format.space_before = Pt(1)
            p_skill.paragraph_format.space_after = Pt(2)
            run_cat = p_skill.add_run(f"{cat.name}: ")
            run_cat.bold = True
            p_skill.add_run(cat.skills)

    # -----------------------------------------------------------------
    # 4. EXPERIENCIA PROFESIONAL (FÓRMULA XYZ & VIÑETAS SÓLIDAS)
    # -----------------------------------------------------------------
    if cv.experiences:
        add_section_header("Experiencia Profesional")
        for exp in cv.experiences:
            # Línea de Cargo y Fechas
            p_exp = doc.add_paragraph()
            p_exp.paragraph_format.space_before = Pt(6)
            p_exp.paragraph_format.space_after = Pt(1)
            p_exp.paragraph_format.keep_with_next = True

            run_role = p_exp.add_run(exp.job_title)
            run_role.bold = True
            run_role.font.size = Pt(10.5)

            end_date = exp.end_date if exp.end_date else "Presente"
            dates_str = f"  |  {exp.start_date} – {end_date}"
            run_dates = p_exp.add_run(dates_str)
            run_dates.font.italic = True
            run_dates.font.size = Pt(9.5)

            # Sub-línea: Empresa y Ubicación
            p_sub = doc.add_paragraph()
            p_sub.paragraph_format.space_before = Pt(0)
            p_sub.paragraph_format.space_after = Pt(3)
            p_sub.paragraph_format.keep_with_next = True

            run_comp = p_sub.add_run(exp.company)
            run_comp.bold = True
            if exp.location:
                p_sub.add_run(f"  •  {exp.location}")

            # Viñetas sólidas de logros cuantificables
            for bullet in exp.bullets:
                p_bullet = doc.add_paragraph(style='List Bullet')
                p_bullet.paragraph_format.space_before = Pt(1)
                p_bullet.paragraph_format.space_after = Pt(2)
                p_bullet.paragraph_format.line_spacing = 1.15
                run_b = p_bullet.add_run(bullet)
                run_b.font.size = Pt(9.8)

    # -----------------------------------------------------------------
    # 5. PROYECTOS DESTACADOS
    # -----------------------------------------------------------------
    if cv.projects:
        add_section_header("Proyectos Destacados")
        for proj in cv.projects:
            p_proj = doc.add_paragraph()
            p_proj.paragraph_format.space_before = Pt(5)
            p_proj.paragraph_format.space_after = Pt(1)
            p_proj.paragraph_format.keep_with_next = True

            run_pname = p_proj.add_run(proj.name)
            run_pname.bold = True
            run_pname.font.size = Pt(10.5)

            if proj.technologies:
                run_tech = p_proj.add_run(f"  |  {proj.technologies}")
                run_tech.font.italic = True
                run_tech.font.size = Pt(9.5)

            if proj.project_url:
                run_url = p_proj.add_run(f"  ({proj.project_url})")
                run_url.font.size = Pt(9)
                run_url.font.color.rgb = RGBColor(37, 99, 235)

            p_pdesc = doc.add_paragraph()
            p_pdesc.paragraph_format.space_before = Pt(0)
            p_pdesc.paragraph_format.space_after = Pt(4)
            p_pdesc.paragraph_format.line_spacing = 1.15
            p_pdesc.add_run(proj.description)

    # -----------------------------------------------------------------
    # 6. EDUCACIÓN Y FORMACIÓN
    # -----------------------------------------------------------------
    if cv.education:
        add_section_header("Educación y Formación Continua")
        for edu in cv.education:
            p_edu = doc.add_paragraph()
            p_edu.paragraph_format.space_before = Pt(4)
            p_edu.paragraph_format.space_after = Pt(2)
            run_deg = p_edu.add_run(edu.degree)
            run_deg.bold = True
            p_edu.add_run(f"  •  {edu.institution}")
            if edu.graduation_year:
                run_yr = p_edu.add_run(f"  ({edu.graduation_year})")
                run_yr.font.italic = True

    # Guardar archivo
    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    doc.save(output_path)
    return output_path
